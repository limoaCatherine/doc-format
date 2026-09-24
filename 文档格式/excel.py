# -*- coding: utf-8 -*-
"""Excel (.xlsx) 轻量排版。"""
from __future__ import annotations

import os
import warnings

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from 文档格式.config import COLOR_HEADER_BG, FONT_NAME, XLSX_FONT_ONLY_KEYWORDS
from 文档格式.scan import atomic_save_xlsx


def _is_cjk(ch: str) -> bool:
    if not ch:
        return False
    cp = ord(ch[0])
    return 0x4E00 <= cp <= 0x9FFF or 0x3000 <= cp <= 0x303F or 0xFF00 <= cp <= 0xFFEF


def _text_width(s: str) -> int:
    return sum(2 if _is_cjk(ch) else 1 for ch in s)


def xlsx_font_only(path: str) -> bool:
    name = os.path.basename(path)
    return any(k in name for k in XLSX_FONT_ONLY_KEYWORDS)


def _set_cell_yahei(cell) -> None:
    old = cell.font
    if old and old.name == FONT_NAME:
        return
    if old:
        cell.font = Font(
            name=FONT_NAME,
            size=old.size,
            bold=old.bold,
            italic=old.italic,
            underline=old.underline,
            strike=old.strike,
            color=old.color,
        )
    else:
        cell.font = Font(name=FONT_NAME)


def _header_has_fill(cell) -> bool:
    fill = cell.fill
    if not fill or fill.fill_type in (None, "none"):
        return False
    fg = getattr(fill, "fgColor", None)
    if fg is None:
        return False
    rgb = getattr(fg, "rgb", None) or getattr(fg, "theme", None)
    if rgb in (None, "00000000", "FFFFFFFF", "00FFFFFF"):
        return False
    return True


def process_xlsx(path: str) -> int:
    font_only = xlsx_font_only(path)
    try:
        fsize = os.path.getsize(path)
    except OSError:
        fsize = 0
    if fsize >= 1_500_000:
        font_only = True

    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")
        wb = openpyxl.load_workbook(path)

    changed = 0
    header_fill = PatternFill("solid", fgColor=COLOR_HEADER_BG)
    tab_color = "B8CCE4"

    for ws in wb.worksheets:
        for cell in ws._cells.values():
            old = cell.font
            old_name = old.name if old else None
            if old_name == FONT_NAME:
                continue
            if cell.value is None and (old is None or old_name in (None, "Calibri", "Arial")):
                if old_name in ("Calibri", "Arial"):
                    _set_cell_yahei(cell)
                    changed += 1
                continue
            _set_cell_yahei(cell)
            changed += 1

        if font_only:
            continue

        max_r = ws.max_row or 0
        max_c = ws.max_column or 0
        if max_r <= 0 or max_c <= 0:
            continue

        row1 = [ws.cell(1, c) for c in range(1, max_c + 1)]
        if any(c.value is not None for c in row1) and not any(
            _header_has_fill(c) for c in row1 if c.value is not None
        ):
            for c in row1:
                if c.value is None:
                    continue
                c.fill = header_fill
                _set_cell_yahei(c)
                old = c.font
                c.font = Font(
                    name=FONT_NAME,
                    size=old.size if old else 11,
                    bold=True,
                    italic=old.italic if old else None,
                    color=old.color if old else None,
                )
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        sample_rows = min(max_r, 80)
        for col_idx in range(1, max_c + 1):
            max_w = 0
            for row_idx in range(1, sample_rows + 1):
                v = ws.cell(row=row_idx, column=col_idx).value
                if v is not None:
                    max_w = max(max_w, _text_width(str(v)[:80]))
            if max_w:
                ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_w + 2, 8), 45)

        try:
            ws.sheet_properties.tabColor = tab_color
        except Exception:
            pass

    atomic_save_xlsx(wb, path)
    wb.close()
    return changed
