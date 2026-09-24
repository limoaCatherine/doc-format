# -*- coding: utf-8 -*-
"""Word (.docx) 排版。"""
from __future__ import annotations

from pathlib import Path
from typing import Optional

from docx import Document as _DocFactory
from docx.document import Document as DocumentType
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table
from docx.text.run import Run

from 文档格式.config import (
    COLOR_BORDER,
    COLOR_H1,
    COLOR_H2,
    COLOR_H3,
    COLOR_HEADER_BG,
    COLOR_HEADER_FG,
    COLOR_TOTAL,
    COLOR_ZEBRA,
    FONT_NAME,
    HIGHLIGHT_KEYWORDS,
    SIZE_BODY,
    SIZE_H1,
    SIZE_H2,
    SIZE_H3,
    SIZE_HEADER,
)
from 文档格式.scan import atomic_save_docx


def set_run_font(
    run: Run,
    size: Optional[Pt] = None,
    bold: Optional[bool] = None,
    color: Optional[RGBColor] = None,
    name: str = FONT_NAME,
) -> None:
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_all_fonts(doc: DocumentType) -> None:
    for p in doc.paragraphs:
        for r in p.runs:
            set_run_font(r)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        set_run_font(r)


def _set_cell_shading(cell, fill_hex: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    tcPr.append(shd)


def _set_cell_valign_center(cell) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:vAlign")):
        tcPr.remove(old)
    vAlign = OxmlElement("w:vAlign")
    vAlign.set(qn("w:val"), "center")
    tcPr.append(vAlign)


def set_table_borders(table: Table) -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    style_elem = tblPr.find(qn("w:tblStyle"))
    if style_elem is not None:
        tblPr.remove(style_elem)

    borders = OxmlElement("w:tblBorders")
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{name}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), COLOR_BORDER)
        borders.append(b)
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    tblPr.append(borders)


def set_header_repeat(table: Table) -> None:
    if not table.rows:
        return
    tr = table.rows[0]._tr
    trPr = tr.get_or_add_trPr()
    for old in trPr.findall(qn("w:tblHeader")):
        trPr.remove(old)
    hdr = OxmlElement("w:tblHeader")
    hdr.set(qn("w:val"), "true")
    trPr.append(hdr)


def set_row_no_split(table: Table) -> None:
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        for old in trPr.findall(qn("w:cantSplit")):
            trPr.remove(old)
        cs = OxmlElement("w:cantSplit")
        cs.set(qn("w:val"), "true")
        trPr.append(cs)


def autofit_table(table: Table) -> None:
    tblPr = table._tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        table._tbl.insert(0, tblPr)
    for old in tblPr.findall(qn("w:tblW")):
        tblPr.remove(old)
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), "5000")
    tblW.set(qn("w:type"), "pct")
    tblPr.append(tblW)
    for old in tblPr.findall(qn("w:tblLayout")):
        tblPr.remove(old)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "autofit")
    tblPr.append(layout)


def is_column_numeric(table: Table, col_idx: int) -> bool:
    numeric = total = 0
    for row in table.rows[1:]:
        if col_idx >= len(row.cells):
            continue
        text = row.cells[col_idx].text.strip()
        total += 1
        if not text:
            continue
        clean = (
            text.replace("%", "")
            .replace(",", "")
            .replace("(", "")
            .replace(")", "")
            .replace("万", "")
            .strip()
        )
        if clean and (clean.replace(".", "", 1).replace("-", "", 1).isdigit() or clean == "-"):
            numeric += 1
    return total > 0 and numeric >= total * 0.8


def format_table(table: Table) -> None:
    """边框 + 表头 + 斑马 + 合计行 + 对齐 + 跨页表头。"""
    if not table.rows:
        return
    set_table_borders(table)
    set_header_repeat(table)
    set_row_no_split(table)
    autofit_table(table)

    n_cols = len(table.rows[0].cells)
    numeric_cols = {c for c in range(n_cols) if is_column_numeric(table, c)}

    for row_idx, row in enumerate(table.rows):
        row_text = " ".join(c.text for c in row.cells)
        is_header = row_idx == 0
        is_total = (not is_header) and any(kw in row_text for kw in HIGHLIGHT_KEYWORDS)
        if is_header:
            fill = COLOR_HEADER_BG
        elif is_total:
            fill = COLOR_TOTAL
        elif row_idx % 2 == 0:
            fill = COLOR_ZEBRA
        else:
            fill = "FFFFFF"

        for col_idx, cell in enumerate(row.cells):
            _set_cell_shading(cell, fill)
            _set_cell_valign_center(cell)
            for p in cell.paragraphs:
                if is_header:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif col_idx in numeric_cols:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    set_run_font(
                        r,
                        size=SIZE_HEADER if is_header else SIZE_BODY,
                        bold=True if (is_header or is_total) else None,
                        color=COLOR_HEADER_FG if is_header else None,
                    )


def set_page_margins(doc: DocumentType) -> None:
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.17)
        section.right_margin = Cm(3.17)
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)


def set_heading_styles(doc: DocumentType) -> None:
    mapping = [
        (("Heading 1", "标题 1"), SIZE_H1, COLOR_H1, Pt(12), Pt(6)),
        (("Heading 2", "标题 2"), SIZE_H2, COLOR_H2, Pt(10), Pt(4)),
        (("Heading 3", "标题 3"), SIZE_H3, COLOR_H3, Pt(8), Pt(3)),
    ]
    for p in doc.paragraphs:
        style_name = (p.style.name if p.style else "") or ""
        for keys, size, color, before, after in mapping:
            if any(k in style_name for k in keys):
                pf = p.paragraph_format
                pf.space_before = before
                pf.space_after = after
                pf.line_spacing = 1.3
                pPr = p._element.get_or_add_pPr()
                for old in pPr.findall(qn("w:pageBreakBefore")):
                    pPr.remove(old)
                for r in p.runs:
                    set_run_font(r, size=size, bold=True, color=color)
                break


def beautify_paragraphs(doc: DocumentType) -> None:
    for p in doc.paragraphs:
        style_name = (p.style.name if p.style else "") or ""
        if "Normal" in style_name or "正文" in style_name:
            pf = p.paragraph_format
            pf.first_line_indent = Cm(0.74)
            pf.line_spacing = 1.5
            pf.space_after = Pt(6)
            for r in p.runs:
                if r.font.size is None:
                    set_run_font(r, size=SIZE_BODY)


def collapse_empty_paragraphs(doc: DocumentType) -> None:
    prev_empty = False
    for p in doc.paragraphs:
        empty = not (p.text or "").strip()
        if empty and prev_empty:
            for r in p.runs:
                r.text = ""
        prev_empty = empty


def center_images(doc: DocumentType) -> None:
    for p in doc.paragraphs:
        for run in p.runs:
            if run._element.findall(qn("w:drawing")):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                break


def _add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    set_run_font(run, size=Pt(9), color=RGBColor(0x66, 0x66, 0x66))


def set_header_footer(doc: DocumentType, title: str) -> None:
    for section in doc.sections:
        header = section.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        hp.text = ""
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = hp.add_run(title)
        set_run_font(run, size=Pt(9), color=RGBColor(0x88, 0x88, 0x88))

        footer = section.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        fp.text = ""
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_page_number(fp)


def process_docx(path: str) -> int:
    doc = _DocFactory(path)
    set_all_fonts(doc)
    set_page_margins(doc)
    set_heading_styles(doc)
    beautify_paragraphs(doc)
    collapse_empty_paragraphs(doc)
    center_images(doc)
    set_header_footer(doc, Path(path).name)
    for t in doc.tables:
        format_table(t)
    atomic_save_docx(doc, path)
    return len(doc.tables)
